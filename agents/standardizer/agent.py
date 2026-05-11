from __future__ import annotations

import hashlib
import json
import logging
import mimetypes
import re
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, AsyncGenerator
from urllib.parse import urlparse

from google.adk.agents import BaseAgent
from google.adk.agents.invocation_context import InvocationContext
from google.adk.events import Event
from google.genai import types

import pipeline_client

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from PIL import Image  # noqa: F401
except Exception:
    Image = None

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"
ARTIFACTS_DIR = DATA_DIR / "artifacts"
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml",
    ".yaml", ".yml", ".log", ".rtf",
}

logger = logging.getLogger(__name__)


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _ensure_storage() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)


def _ensure_schema(connection: object | None = None) -> None:
    # Schema is managed by Django migrations; nothing to do at runtime.
    pass


def _to_iso_or_none(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), tzinfo=timezone.utc).isoformat()
    text = str(value).strip()
    return text or None


def _normalize_text(text: str | None) -> str | None:
    if not text:
        return None
    normalized = re.sub(r"\s+", " ", text).strip()
    return normalized or None


def _safe_json(value: Any) -> Any:
    """Return value as a JSON-safe structure (for JSONField storage)."""
    try:
        return json.loads(json.dumps(value, ensure_ascii=True, default=str))
    except Exception:
        return {}


def _load_screened_rows_from_db() -> list[dict[str, Any]]:
    try:
        rows = pipeline_client.load_screened_files_for_standardization()
        logger.debug("Loaded %d selected unprocessed screener rows", len(rows))
        return rows
    except Exception:
        logger.exception("Failed loading screener rows from API")
        return []


def _dedupe_selected_screened_rows(
    items: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[int]]:
    latest_by_path: dict[str, dict[str, Any]] = {}
    duplicate_screened_file_ids: list[int] = []

    for item in items:
        source_path = str(item.get("path") or item.get("pageUrl") or "").strip()
        dedupe_key = source_path.lower() if source_path else f"screened_file_id:{item.get('_screened_file_id')}"
        existing = latest_by_path.get(dedupe_key)

        if existing is None:
            latest_by_path[dedupe_key] = item
            continue

        current_id = item.get("_screened_file_id")
        existing_id = existing.get("_screened_file_id")
        if isinstance(current_id, int) and isinstance(existing_id, int) and current_id > existing_id:
            duplicate_screened_file_ids.append(existing_id)
            latest_by_path[dedupe_key] = item
        elif isinstance(current_id, int):
            duplicate_screened_file_ids.append(current_id)

    deduped_items = sorted(
        latest_by_path.values(),
        key=lambda i: int(i.get("_screened_file_id")) if isinstance(i.get("_screened_file_id"), int) else 0,
    )
    return deduped_items, duplicate_screened_file_ids


def _load_screened_payload(ctx: InvocationContext) -> tuple[list[dict[str, Any]], str]:
    del ctx
    from_db = _load_screened_rows_from_db()
    return from_db, "postgresql.screened_files"



def _sha256_hex(text: str | None) -> str | None:
    if not text:
        return None
    return hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()


def _guess_mime(path_or_url: str) -> str:
    mime, _ = mimetypes.guess_type(path_or_url)
    return mime or "application/octet-stream"


def _file_metadata(path: Path) -> dict[str, Any]:
    if not path.exists() or not path.is_file():
        return {"exists": False, "suffix": path.suffix.lower(), "name": path.name, "path": str(path)}

    stat = path.stat()
    return {
        "exists": True,
        "name": path.name,
        "path": str(path),
        "suffix": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "mime_type": _guess_mime(str(path)),
        "created_at": datetime.fromtimestamp(stat.st_ctime, tz=timezone.utc).isoformat(),
        "modified_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
    }


def _read_text_file(path: Path) -> tuple[str | None, str | None]:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return path.read_text(encoding=encoding), None
        except UnicodeDecodeError:
            continue
        except Exception as error:
            logger.exception("Text extraction failed for %s", path)
            return None, f"Text extraction failed: {error}"
    return None, "Text extraction failed: unsupported character encoding"


def _extract_html(path: Path) -> tuple[str | None, dict[str, Any], list[dict[str, str]], str | None]:
    if BeautifulSoup is None:
        return None, {}, [], "Missing dependency: beautifulsoup4"

    try:
        html = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        html = path.read_text(encoding="latin-1", errors="replace")
    except Exception as error:
        logger.exception("HTML extraction failed for %s", path)
        return None, {}, [], f"HTML extraction failed: {error}"

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "template", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)

    def _meta_value(*keys: str) -> str | None:
        lowered = {k.lower() for k in keys}
        for meta in soup.find_all("meta"):
            for attr in ("name", "property", "itemprop", "http-equiv"):
                raw = meta.attrs.get(attr)
                if isinstance(raw, str) and raw.lower() in lowered:
                    content = meta.attrs.get("content")
                    if isinstance(content, str) and content.strip():
                        return content.strip()
        return None

    title_text: str | None = None
    if soup.title and soup.title.string:
        title_text = soup.title.string.strip() or None

    published = _meta_value("article:published_time", "publish_date", "date", "pubdate", "dc.date", "datepublished")
    modified = _meta_value("article:modified_time", "last-modified", "datemodified", "dc.modified")
    author = _meta_value("author", "article:author", "dc.creator")

    media: list[dict[str, str]] = []
    for img in soup.find_all("img"):
        src = img.attrs.get("src")
        if not isinstance(src, str) or not src.strip():
            continue
        src = src.strip()
        parsed = urlparse(src)
        resolved = src if parsed.scheme in {"http", "https"} else str((path.parent / src).resolve())
        media.append({"artifact_path": resolved, "mime_type": _guess_mime(resolved)})

    metadata = {
        "title": title_text or _meta_value("og:title", "twitter:title"),
        "author": author,
        "description": _meta_value("description", "og:description", "twitter:description"),
        "published_at": published,
        "modified_at": modified,
        "language": _meta_value("og:locale", "language", "dc.language"),
        "source_type": "html",
    }
    return _normalize_text(text), metadata, media, None


def _extract_pdf(path: Path) -> tuple[str | None, dict[str, Any], list[dict[str, str]], str | None]:
    if PdfReader is None:
        return None, {}, [], "Missing dependency: pypdf"

    try:
        reader = PdfReader(str(path))
    except Exception as error:
        logger.exception("PDF open failed for %s", path)
        return None, {}, [], f"PDF open failed: {error}"

    page_text: list[str] = []
    media: list[dict[str, str]] = []
    page_errors: list[str] = []
    can_extract_images = Image is not None
    if not can_extract_images:
        page_errors.append("PDF image extraction skipped: Pillow not installed.")

    for index, page in enumerate(reader.pages):
        try:
            extracted = page.extract_text(extraction_mode="layout")
        except TypeError:
            extracted = page.extract_text()
        except Exception as error:
            page_errors.append(f"page {index + 1}: {error}")
            continue
        if extracted:
            page_text.append(extracted)

        if can_extract_images:
            try:
                images = list(page.images)
            except Exception as error:
                page_errors.append(f"page {index + 1} images: {error}")
                images = []

            for image_index, image in enumerate(images):
                try:
                    image_name = str(getattr(image, "name", "") or f"p{index + 1}_img{image_index + 1}.bin")
                    image_data = getattr(image, "data", None)
                    if not isinstance(image_data, (bytes, bytearray)) or not image_data:
                        continue
                    artifact_stem = hashlib.sha256(
                        f"{path}:{index}:{image_index}:{image_name}".encode("utf-8")
                    ).hexdigest()[:16]
                    output_name = f"{path.stem}_{artifact_stem}_{Path(image_name).name}"
                    output_path = ARTIFACTS_DIR / output_name
                    output_path.write_bytes(bytes(image_data))
                    media.append({"artifact_path": str(output_path), "mime_type": _guess_mime(output_name)})
                except Exception as error:
                    page_errors.append(f"page {index + 1} image {image_index + 1}: {error}")

    text = _normalize_text("\n".join(page_text))
    meta = reader.metadata
    metadata = {
        "title": getattr(meta, "title", None) if meta else None,
        "author": getattr(meta, "author", None) if meta else None,
        "subject": getattr(meta, "subject", None) if meta else None,
        "creator": getattr(meta, "creator", None) if meta else None,
        "producer": getattr(meta, "producer", None) if meta else None,
        "creation_date": _to_iso_or_none(getattr(meta, "creation_date", None) if meta else None),
        "modification_date": _to_iso_or_none(getattr(meta, "modification_date", None) if meta else None),
        "is_encrypted": bool(getattr(reader, "is_encrypted", False)),
        "page_count": len(reader.pages),
        "source_type": "pdf",
    }
    error: str | None = "; ".join(page_errors) if page_errors else None
    if not text and not error:
        error = "No extractable text found (possibly scanned/image-only PDF)"
    return text, metadata, media, error


def _extract_docx_media(path: Path) -> list[dict[str, str]]:
    media_refs: list[dict[str, str]] = []
    artifact_prefix = hashlib.sha256(str(path).encode("utf-8")).hexdigest()[:12]
    try:
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if not member.startswith("word/media/"):
                    continue
                filename = Path(member).name
                if not filename:
                    continue
                data = archive.read(member)
                output_name = f"{artifact_prefix}_{filename}"
                output_path = ARTIFACTS_DIR / output_name
                output_path.write_bytes(data)
                media_refs.append({"artifact_path": str(output_path), "mime_type": _guess_mime(output_name)})
    except Exception:
        logger.exception("Failed extracting DOCX media from %s", path)
        return []
    return media_refs


def _extract_docx(path: Path) -> tuple[str | None, dict[str, Any], list[dict[str, str]], str | None]:
    if Document is None:
        return None, {}, [], "Missing dependency: python-docx"

    try:
        document = Document(str(path))
    except Exception as error:
        logger.exception("Word document open failed for %s", path)
        return None, {}, [], f"Word document open failed: {error}"

    chunks: list[str] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            chunks.append(value)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    core = document.core_properties
    metadata = {
        "title": core.title or None,
        "author": core.author or None,
        "subject": core.subject or None,
        "keywords": core.keywords or None,
        "category": core.category or None,
        "comments": core.comments or None,
        "last_modified_by": core.last_modified_by or None,
        "created": _to_iso_or_none(core.created),
        "modified": _to_iso_or_none(core.modified),
        "source_type": "word",
        "word_format": path.suffix.lower(),
    }
    media = _extract_docx_media(path)
    text = _normalize_text("\n".join(chunks))
    error = None if text else "No extractable text found in Word document"
    return text, metadata, media, error


def _extract_media_refs_from_item(item: dict[str, Any]) -> list[dict[str, str]]:
    media_refs: list[dict[str, str]] = []
    for key in ("media", "mediaPaths", "attachments"):
        value = item.get(key)
        if isinstance(value, list):
            for candidate in value:
                if isinstance(candidate, str) and candidate.strip():
                    c = candidate.strip()
                    media_refs.append({"artifact_path": c, "mime_type": _guess_mime(c)})
    return media_refs


def _infer_author(text: str | None, metadata_author: str | None, source_author: str | None) -> str | None:
    for candidate in (source_author, metadata_author):
        if isinstance(candidate, str) and candidate.strip():
            return candidate.strip()

    if not text:
        return None

    lines = [line.strip() for line in text.splitlines() if line.strip()][:20]
    patterns = [
        re.compile(r"^(?:by|author)\s*[:\-]\s*(.+)$", re.IGNORECASE),
        re.compile(r"^by\s+([A-Z][A-Za-z.'\- ]{1,120})$"),
    ]
    for line in lines:
        for pattern in patterns:
            match = pattern.match(line)
            if match:
                value = match.group(1).strip()
                if value:
                    return value
    return None


def _extract_from_file(path: Path) -> tuple[str | None, dict[str, Any], list[dict[str, str]], str | None]:
    suffix = path.suffix.lower()
    if not path.exists() or not path.is_file():
        return None, {}, [], f"File not found: {path}"
    if suffix in {".html", ".htm"}:
        return _extract_html(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".docx", ".docm"}:
        return _extract_docx(path)
    if suffix == ".doc":
        return None, {"source_type": "word", "word_format": ".doc"}, [], (
            "Legacy .doc is unsupported; convert to .docx"
        )
    if suffix in TEXT_EXTENSIONS:
        text, error = _read_text_file(path)
        return _normalize_text(text), {"source_type": "plain_text"}, [], error
    return None, {"source_type": "unknown"}, [], f"Unsupported file type: {suffix or '[no extension]'}"


def _is_web_source_item(item: dict[str, Any]) -> bool:
    path = str(item.get("path") or item.get("pageUrl") or "").strip().lower()
    if path.startswith("http://") or path.startswith("https://"):
        return True
    if str(item.get("sourceType") or "").strip().lower() == "webpage":
        return True
    return any(key in item for key in ("contentText", "pageSummary", "newsItems"))


def _web_source_metadata(item: dict[str, Any]) -> dict[str, Any]:
    path = str(item.get("path") or item.get("pageUrl") or "").strip()
    name = str(item.get("name") or item.get("pageTitle") or path).strip()
    return {
        "exists": False,
        "name": name or path,
        "path": path,
        "suffix": ".html",
        "mime_type": "text/html",
        "source_type": "webpage",
    }


def _extract_from_web_item(item: dict[str, Any]) -> tuple[str | None, dict[str, Any], list[dict[str, str]], str | None]:
    news_items = item.get("newsItems")
    news_lines: list[str] = []
    if isinstance(news_items, list):
        for news in news_items:
            if not isinstance(news, dict):
                continue
            parts = [
                p for p in [
                    str(news.get("title") or "").strip(),
                    str(news.get("summary") or "").strip(),
                    str(news.get("publishedAt") or "").strip(),
                    str(news.get("url") or "").strip(),
                ]
                if p
            ]
            if parts:
                news_lines.append(" | ".join(parts))

    text_parts = [
        str(item.get("pageTitle") or "").strip(),
        str(item.get("pageSummary") or "").strip(),
        str(item.get("contentText") or "").strip(),
        "\n".join(news_lines).strip(),
    ]
    text_content = _normalize_text("\n\n".join(part for part in text_parts if part))
    metadata = {
        "title": str(item.get("pageTitle") or item.get("name") or "").strip() or None,
        "description": str(item.get("pageSummary") or "").strip() or None,
        "published_at": str(item.get("createdAt") or "").strip() or None,
        "modified_at": str(item.get("modifiedAt") or "").strip() or None,
        "company": str(item.get("company") or "").strip() or None,
        "news_items": news_items if isinstance(news_items, list) else [],
        "source_type": "webpage",
    }
    error = str(item.get("scrapeError") or "").strip() or None
    return text_content, metadata, [], error


class DocumentStandardizerAgent(BaseAgent):
    def __init__(self) -> None:
        super().__init__(
            name="document_standardizer_agent",
            description="Reads screener output, extracts document content and metadata, and persists standardized records.",
        )

    async def _run_async_impl(self, ctx: InvocationContext) -> AsyncGenerator[Event, None]:
        _ensure_storage()
        logger.debug("Starting standardizer run")

        screened_items, screened_source = _load_screened_payload(ctx)
        logger.debug("Loaded %d screened items from %s", len(screened_items), screened_source)

        selected_items = [item for item in screened_items if bool(item.get("isSelected", True))]
        selected_items, duplicate_screened_file_ids = _dedupe_selected_screened_rows(selected_items)
        logger.debug(
            "Selected %d deduped items. skipped_duplicates=%d",
            len(selected_items),
            len(duplicate_screened_file_ids),
        )
        standardized_documents: list[dict[str, Any]] = []

        for screened_file_id in duplicate_screened_file_ids:
            pipeline_client.mark_screened_file_processed(screened_file_id)
            logger.debug("Marked duplicate screened_file_id=%s processed", screened_file_id)

        for item in selected_items:
            source_path = str(item.get("path") or "").strip()
            source_name = str(item.get("name") or "").strip()
            is_web_item = _is_web_source_item(item)

            if is_web_item:
                resolved_path_str = source_path or str(item.get("pageUrl") or "").strip() or source_name
                text_content, extracted_metadata, extracted_media, extraction_error = _extract_from_web_item(item)
                filesystem_meta = _web_source_metadata(item)
            else:
                candidate = Path(source_path) if source_path else Path(source_name)
                resolved_path = candidate.expanduser().resolve()
                resolved_path_str = str(resolved_path)
                text_content, extracted_metadata, extracted_media, extraction_error = _extract_from_file(resolved_path)
                filesystem_meta = _file_metadata(resolved_path)

            status = "ok" if extraction_error is None else "error"
            item_media = _extract_media_refs_from_item(item)
            media_items = extracted_media + item_media
            author = _infer_author(
                text_content,
                extracted_metadata.get("author") if isinstance(extracted_metadata, dict) else None,
                item.get("author") if isinstance(item.get("author"), str) else None,
            )
            metadata = {
                "source": item,
                "filesystem": filesystem_meta,
                "extracted": extracted_metadata,
            }

            content_sha256 = _sha256_hex(text_content)
            screened_file_id = item.get("_screened_file_id")
            hoarder_output_id = item.get("_hoarder_output_id")

            result = pipeline_client.standardize_document(
                source_path=resolved_path_str,
                author=author,
                text_content=text_content,
                metadata_json=_safe_json(metadata),
                extraction_status=status,
                extraction_error=extraction_error,
                content_sha256=content_sha256,
                modified_at=item.get("modifiedAt"),
                created_at=item.get("createdAt"),
                screened_file_id=screened_file_id if isinstance(screened_file_id, int) else None,
                hoarder_output_id=hoarder_output_id if isinstance(hoarder_output_id, int) else None,
                media=media_items,
            )
            doc_id = result.get("doc_id")
            logger.debug("Inserted document doc_id=%s source_path=%s", doc_id, resolved_path_str)

            persisted_media: list[dict[str, str]] = []
            for media_ref in media_items:
                artifact_path = str(media_ref.get("artifact_path") or "").strip()
                if not artifact_path:
                    continue
                mime_type = str(media_ref.get("mime_type") or _guess_mime(artifact_path))
                persisted_media.append({"artifact_path": artifact_path, "mime_type": mime_type})

            standardized_documents.append({
                "doc_id": doc_id,
                "text": text_content,
                "author": author,
                "metadata": metadata,
                "media": persisted_media,
                "extraction": {"status": status, "error": extraction_error},
                "content_sha256": content_sha256,
                "screened_file_id": screened_file_id,
            })

        standardizer_payload = {
            "persistedCount": len(standardized_documents),
            "screenedInputSource": screened_source,
            "documents": standardized_documents,
        }
        ctx.session.state["standardized_documents"] = json.dumps(standardized_documents, default=str)
        logger.debug("Standardizer committed %d documents", len(standardized_documents))

        yield Event(
            author=self.name,
            invocation_id=ctx.invocation_id,
            content=types.Content(
                role="model",
                parts=[types.Part(text=json.dumps(standardizer_payload, indent=2, default=str))],
            ),
        )


standardizer_agent = DocumentStandardizerAgent()
